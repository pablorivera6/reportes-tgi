"""Las conclusiones y recomendaciones no se pierden por falta de espacio.

La hoja Informe deja pocas filas para cada bloque (CIPS: 8 conclusiones y 2
recomendaciones; PAP: 4 recomendaciones) y la base automática genera más. Lo que
no cabía quedaba FUERA del informe con el aviso "N conclusión(es)/
recomendación(es) no cupieron" (Tausa, 2026-10). Ahora se abren filas: lo de
abajo (recomendaciones, firmas) baja y las filas nuevas copian el formato."""
import os
import re

import openpyxl
import pytest

from generator import ReportGenerator, resource_path

PLANTILLAS = ["EN BLANCO.xlsx", "CIPS EN BLANCO.xlsx", "DCVG_REP.xlsx"]
CONC = [f"Conclusión número {i} del tramo inspeccionado." for i in range(1, 19)]
RECO = [f"Recomendación número {i} para el tramo." for i in range(1, 17)]
FIRMAS = {'nombre': 'Javier Jara', 'cargo': 'Ingeniero Residente', 'empresa': 'PCC Integrity'}


def _generar(tmp_path, plantilla, conc=CONC, reco=RECO):
    gen = ReportGenerator(resource_path(plantilla))
    gen.fill_conclusiones(conc)
    gen.fill_recomendaciones(reco)
    gen.fill_firmas({'nombre': 'Pablo Rivera', 'cargo': 'Ingeniero Junior',
                     'empresa': 'PCC Integrity'}, FIRMAS,
                    {'nombre': 'Alejandro Rivera', 'cargo': 'Ingeniero Especialista CP4',
                     'empresa': 'PCC Integrity'})
    out = os.path.join(tmp_path, "x.xlsx")
    gen.save(out)
    return gen, openpyxl.load_workbook(out)


@pytest.fixture(scope="module", params=PLANTILLAS)
def generado(request, tmp_path_factory):
    """Un informe por plantilla con 18 conclusiones y 16 recomendaciones
    (se genera una vez: guardar la plantilla CIPS tarda)."""
    return _generar(tmp_path_factory.mktemp("conc"), request.param)


def _textos(ws):
    return {str(c.value).strip(): c.row for fila in ws.iter_rows(max_col=1)
            for c in fila if isinstance(c.value, str)}


def test_ninguna_queda_por_fuera(generado):
    gen, wb = generado
    assert gen.conclusiones_omitidas == 0 and gen.recomendaciones_omitidas == 0
    filas = _textos(wb["Informe"])
    c = [filas[f"• {t}"] for t in CONC]
    r = [filas[f"• {t}"] for t in RECO]
    assert c == sorted(c) and r == sorted(r)
    titulo_reco = next(f for t, f in filas.items() if t.upper().startswith("RECOMENDACIONES"))
    assert c[-1] < titulo_reco < r[0]


def test_las_firmas_bajan_y_las_demas_hojas_las_siguen(generado):
    gen, wb = generado
    ws = wb["Informe"]
    filas = _textos(ws)
    ultima_reco = filas[f"• {RECO[-1]}"]
    bloque = gen._bloque_firmas(ws)
    assert min(r for r, _ in bloque.values()) > ultima_reco
    r, c = bloque["REVISÓ"]
    assert ws.cell(r, c).value == "Javier Jara"
    # el área de impresión llega hasta las firmas
    ultima_fila = int(re.findall(r"\d+", str(ws.print_area))[-1])
    assert ultima_fila >= r + 2
    # otra hoja con firmas por fórmula apunta a la celda real
    hoja = wb["Hallazgos"]
    b2 = gen._bloque_firmas(hoja, desde=12)
    r2, c2 = b2["REVISÓ"]
    assert hoja.cell(r2, c2).value.endswith(f"{openpyxl.utils.get_column_letter(c)}{r}")


def test_filas_nuevas_con_el_formato_del_bloque(generado):
    gen, wb = generado
    ws = wb["Informe"]
    fila = _textos(ws)[f"• {RECO[-1]}"]
    combinada = [m for m in ws.merged_cells.ranges if m.min_row == fila == m.max_row]
    assert combinada and combinada[0].min_col == 1 and combinada[0].max_col > 20
    assert ws.cell(fila, 1).alignment.wrap_text


def test_parrafo_largo_no_queda_cortado(tmp_path):
    largo = "Se recomienda" + " verificar el estado del recubrimiento" * 12
    gen, wb = _generar(tmp_path, "EN BLANCO.xlsx", conc=[largo], reco=[])
    ws = wb["Informe"]
    fila = _textos(ws)[f"• {largo}"]
    assert ws.row_dimensions[fila].height >= 40


def test_sin_exceso_la_plantilla_no_cambia(tmp_path):
    """Con pocos párrafos nada se mueve: las firmas siguen donde las trae la plantilla."""
    gen, wb = _generar(tmp_path, "EN BLANCO.xlsx", conc=CONC[:2], reco=RECO[:1])
    assert min(r for r, _ in gen._bloque_firmas(wb["Informe"]).values()) == 100
