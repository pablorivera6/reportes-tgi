"""SISTEMA INSPECCIONADO y MONITOREO de la hoja Informe, por ETIQUETA.

`fill_sistema_inspeccionado` y `fill_monitoreo` tenían quemadas las filas de la
plantilla PAP (39-46 y 50-51). La plantilla CIPS tiene todo UNA fila más
arriba, así que en un CIPS escribían el tipo de inspección sobre AMENAZA,
'Normal' sobre TIPO DE DUCTO, ceros sobre las fórmulas de longitud y el
criterio sobre la fórmula de Datos/km; además la longitud total quedaba en el
37,5 del ejemplo (Ramal Pradera). En PAP, 'Datos/km' se pisaba con 0.
"""
import os

import openpyxl
import pytest

from generator import ReportGenerator, resource_path

INFO = {'fecha': '12/03/2025', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Salento',
        'tipo_ducto': 'Ramal', 'tipo_recubrimiento': 'FBE', 'diametro': '6',
        'contrato': '551007370', 'ot': '1300012786'}


def _pots(n=10):
    return [{'abscisa': i * 1000, 'abscisa_str': f'{i:03d}+000', 'fecha': '12/03/2025',
             'ref_geografica': 'Poste de Potencial', 'on_mv': -1100,
             'off_mv': -900 if i < 8 else -700, 'lat': 4.6 + i / 100, 'lon': -75.5 - i / 100}
            for i in range(n)]


def _cips(n=100):
    return [{'abscisa_val': i * 100, 'off_mv': -950, 'off_limpio': -950, 'on_mv': -1200,
             'on_limpio': -1200, 'lat': 4.6, 'lon': -75.5} for i in range(n)]


def _pap(tmp_path, info, pots):
    gen = ReportGenerator(resource_path("EN BLANCO.xlsx"))
    gen.fill_general_info(dict(info))
    gen.fill_sistema_inspeccionado(dict(info), pots)
    gen.fill_monitoreo(dict(info), pots)
    out = os.path.join(tmp_path, "pap.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)["Informe"]


def _cips_informe(tmp_path, info, cips, pots=None):
    gen = ReportGenerator(resource_path("CIPS EN BLANCO.xlsx"))
    gen.fill_general_info(dict(info))
    gen.fill_sistema_inspeccionado(dict(info), pots or [], cips=cips)
    gen.fill_monitoreo(dict(info), pots or [], cips=cips)
    out = os.path.join(tmp_path, "cips.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)["Informe"]


# ── PAP ──────────────────────────────────────────────────────────────────────

def test_pap_puntos_longitudes_y_porcentajes(tmp_path):
    ws = _pap(tmp_path, dict(INFO, tipo_inspeccion='PAP', longitud_km=12.3), _pots())
    assert ws["G39"].value == 0 and ws["G40"].value == 9000
    assert abs(ws["O39"].value - 4.6) < 1e-9 and abs(ws["W40"].value + 75.59) < 1e-9
    assert ws["G41"].value == 'Inspección PAP'          # no 'PAP' a secas
    assert ws["O42"].value == 'Ramal'
    assert ws["G44"].value == pytest.approx(12.3)       # total: la del tramo
    assert ws["O44"].value == pytest.approx(9.0)        # inspeccionada: lo recorrido
    assert ws["W44"].value == pytest.approx(7.2)        # 8 de 10 protegidos
    assert ws["AC44"].value == pytest.approx(1.8)
    assert ws["O45"].value == pytest.approx(0.8)
    assert ws["AC45"].value == pytest.approx(0.0)


def test_pap_sin_longitud_del_tramo_usa_el_recorrido(tmp_path):
    ws = _pap(tmp_path, dict(INFO, tipo_inspeccion='PAP'), _pots())
    assert ws["G44"].value == pytest.approx(9.0)


def test_pap_descripcion_de_la_linea_sin_el_ejemplo(tmp_path):
    ws = _pap(tmp_path, dict(INFO, tipo_inspeccion='PAP', longitud_km=12.3), _pots())
    assert "Pereira" not in str(ws["A30"].value)
    assert "Ramal Salento" in str(ws["A30"].value) and "12.3" in str(ws["A30"].value)


def test_pap_monitoreo_conserva_criterio_y_datos_km_es_formula(tmp_path):
    ws = _pap(tmp_path, dict(INFO, tipo_inspeccion='PAP'), _pots())
    assert ws["G50"].value == '6.2.1.3 (-850mVCSE)'
    f = str(ws["G51"].value)
    assert f.startswith("=") and "'Potenciales PAP'!B12:B21" in f and "O44" in f


def test_pap_sin_potenciales_no_rompe(tmp_path):
    ws = _pap(tmp_path, dict(INFO, tipo_inspeccion='PAP'), [])
    assert ws["G41"].value == 'Inspección PAP'


# ── CIPS (una fila más arriba que PAP) ───────────────────────────────────────

def test_cips_no_pisa_amenaza_ni_tipo_de_ducto_ni_aerea(tmp_path):
    ws = _cips_informe(tmp_path, dict(INFO, tipo_inspeccion='CIPS'), _cips())
    assert ws["G40"].value == 'Inspección CIPS'
    assert ws["G41"].value == 'CORROSIÓN EXTERNA'
    assert ws["O41"].value == 'Ramal'
    assert ws["W41"].value == 'Corriente Impresa'
    assert ws["G42"].value == 0 and ws["O42"].value == 0
    # los extremos siguen siendo las fórmulas MIN/MAX sobre Potenciales CIPS
    assert str(ws["G38"].value).startswith("=") and "Potenciales CIPS" in str(ws["G38"].value)
    assert str(ws["O43"].value).startswith("=")           # inspeccionada: fórmula
    for c in ("AC38", "AC39"):                             # altura del ejemplo fuera
        assert "msnm" not in str(ws[c].value or "")


def test_cips_longitud_total_y_formulas_sobre_potenciales_cips(tmp_path):
    ws = _cips_informe(tmp_path, dict(INFO, tipo_inspeccion='CIPS', longitud_km=20.0), _cips())
    assert ws["G43"].value == pytest.approx(20.0)         # no el 37,5 del ejemplo
    g44 = str(ws["G44"].value)
    assert "Potenciales CIPS" in g44 and "F12:F111" in g44 and "Potenciales PAP" not in g44
    # los % se calculan sobre la longitud INSPECCIONADA (O43), no sobre la total
    assert "/O43" in str(ws["O44"].value) and "/O43" in str(ws["W44"].value) \
        and "/O43" in str(ws["AC44"].value)
    assert ws["G49"].value == '6.2.1.3 (-850mVCSE)'
    g50 = str(ws["G50"].value)
    assert "'Potenciales CIPS'!B12:B111" in g50 and "O43" in g50


def test_cips_sin_longitud_del_tramo_usa_el_recorrido(tmp_path):
    ws = _cips_informe(tmp_path, dict(INFO, tipo_inspeccion='CIPS'), _cips())
    assert ws["G43"].value == pytest.approx(9.9)


def test_cips_con_postes_cargados_no_escribe_las_filas_de_pap(tmp_path):
    """Con el FastField PAP de la misma campaña cargado, antes se escribían
    ceros y totales en las filas 43-45 del layout PAP (= 43-45 del CIPS)."""
    ws = _cips_informe(tmp_path, dict(INFO, tipo_inspeccion='CIPS', longitud_km=20.0),
                       _cips(), pots=_pots())
    assert ws["G43"].value == pytest.approx(20.0)
    assert str(ws["O43"].value).startswith("=") and str(ws["W43"].value).startswith("=")
    assert str(ws["G44"].value).startswith("=")
    assert ws["A45"].value and "RESUMEN" in str(ws["A45"].value)
