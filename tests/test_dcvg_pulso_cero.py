"""DCVG: gráfica y observaciones deben contar las MISMAS indicaciones.

Observación del revisor (Ramal Ansermanuevo): la gráfica mostraba 1 defecto y
las observaciones reportaban 4 ("para tres no fue posible calcular el %IR").
Causa: un poste con ON=0 y OFF=0 (fila sin lectura) entraba como poste "con
pulso" (pulso 0) → P/RE = 0 → %IR = #DIV/0! en Excel y None en Python. Los
postes sin pulso real se saltan en la interpolación, y la app avisa de las
indicaciones que de todos modos quedan sin %IR (y por qué)."""
import os

import openpyxl

import db
from dcvg_reader import defectos_sin_severidad
from generator import ReportGenerator, resource_path


def _postes():
    return [{"tipo": "Poste de Potencial", "pk_m": 0, "on": -1200.0, "off": -1000.0},
            {"tipo": "Poste de Potencial", "pk_m": 5000, "on": 0, "off": 0},        # sin lectura
            {"tipo": "Poste de Potencial", "pk_m": 7000, "on": -900.0, "off": -900.0},  # ON = OFF
            {"tipo": "Poste de Potencial", "pk_m": 10000, "on": -1100.0, "off": -900.0}]


def _defectos():
    return [{"pk_m": 4000, "ol_re": 20, "caracter": "AA"},
            {"pk_m": 6000, "ol_re": 50, "caracter": "CA"},
            {"pk_m": 8000, "ol_re": 100, "caracter": "CC"}]


def test_severidad_python_salta_los_postes_sin_pulso():
    sev = db._severidad_dcvg(_postes(), _defectos())
    # pulso 200 en ambos extremos reales → P/RE 200 en todos
    assert [s["p_re"] for s in sev] == [200.0, 200.0, 200.0]
    assert [s["severidad_pct"] for s in sev] == [10.0, 25.0, 50.0]
    assert [s["clasificacion"] for s in sev] == ["Muy Pequeño", "Pequeño", "Mediano"]


def test_formula_excel_interpola_entre_postes_con_pulso_real(tmp_path):
    gen = ReportGenerator(resource_path("DCVG_REP.xlsx"))
    gen.fill_dcvg(_postes(), _defectos())
    out = os.path.join(tmp_path, "d.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)["Inspección DCVG"]
    filas = {ws.cell(r, 4).value: r for r in range(8, 8 + 7)}
    f0, f10 = filas[0], filas[10000]
    for pk in (4000, 6000, 8000):
        q = str(ws.cell(filas[pk], 17).value)
        assert f"P{f0}" in q and f"P{f10}" in q, q
        assert f"P{filas[5000]}" not in q and f"P{filas[7000]}" not in q, q
        # la severidad va escrita (=M/Q) en su columna de carácter
        assert any(str(ws.cell(filas[pk], c).value or "").startswith("=M") for c in (19, 20, 21))


def test_aviso_de_indicaciones_sin_severidad():
    postes = [{"pk_m": 0, "on": 0, "off": 0}]                      # ningún poste con pulso
    defectos = [{"pk_m": 1000, "ol_re": 20, "caracter": "AA"},
                {"pk_m": 2000, "ol_re": None, "caracter": "AA"},
                {"pk_m": 3000, "ol_re": 30, "caracter": ""},
                {"pk_m": None, "ol_re": 30, "caracter": "CC"}]
    avisos = defectos_sin_severidad(postes, defectos)
    motivos = {a for _pk, a in avisos}
    assert len(avisos) == 4
    assert any("pulso" in m.lower() for m in motivos)
    assert any("OL/RE" in m for m in motivos)
    assert any("carácter" in m.lower() for m in motivos)
    assert any("PK" in m or "abscisa" in m.lower() for m in motivos)


def test_sin_problemas_no_hay_aviso():
    assert defectos_sin_severidad(_postes(), _defectos()) == []
