"""Casos reales del FastField DCVG (Ramal Marsella, 2026-07):

- el técnico escribe el PK como '3.000', '6.000.' o '5.00' (kilómetros con
  punto) y el lector solo entendía '3+000': 8 de 9 postes quedaban sin abscisa;
- FastField exporta una fila EN BLANCO en subform_9 / subform_7 cuando ese envío
  no registró defectos o resistividades: salían como defecto/resistividad
  fantasma;
- la columna 'Técnico a cargo' trae un archivo de foto, no un nombre: el
  inspector salía como '258143_….jpg'.
"""
import openpyxl
import pytest

from dcvg_reader import (parse_pk, leer_dcvg_fastfield, leer_resistividades_fastfield,
                         info_desde_meta)


@pytest.mark.parametrize("texto,esperado", [
    ("3+000", 3000), ("PK 7+000", 7000),
    ("3.000", 3000), ("6.000.", 6000), ("12.500", 12500), ("0.000", 0),
    ("5.00", 5000), ("5.5", 5500), ("3,000", 3000), ("2.75", 2750),
    ("2+", None), ("", None), (None, None), ("abc", None), ("150", None),
])
def test_parse_pk_acepta_kilometros_con_punto(texto, esperado):
    assert parse_pk(texto) == esperado


def _root(ws, tecnico):
    ws.append(["FormId", "Form Name", "Submitted On", "Form Version", "Submitted By",
               "Submission Id", "Contratista", "Fecha", "Cliente", "Troncal o ramal inspeccionado",
               "Técnico a cargo"])
    ws.append(["1160295", "Inspección DCVG", "07-06-2026", "26", "x@pcc", "sub-1", "PCC ",
               "07-05-2026", "TGI", "Ramal Marsella", tecnico])


def _dcvg(path, tecnico="258143_315b35e4-6b6c-44f3-804b-d128dbfb2d23_1f7b28eb.jpg"):
    wb = openpyxl.Workbook()
    _root(wb.active, tecnico); wb.active.title = "Root"
    s5 = wb.create_sheet("subform_5")
    s5.append(["Submission Id", "Tipo de poste", "PK", "ON", "Foto ON", "OFF", "Foto OFF",
               "Voltaje AC", "Foto Voltaje AC", "Resistencia ", "Foto resistencia", "Coordenadas"])
    s5.append(["sub-1", "Poste de potencial", "3.000", "-2370", "f", "-1456", "f", "0", "f", "0", "f", "4.95,-75.77"])
    s5.append(["sub-1", "Poste de potencial", "7+000", "-1727", "f", "-1127", "f", "1.1", "f", "0", "f", "4.94,-75.74"])
    s5.append(["sub-2", None, None, None, None, None, None, None, None, None, None, None])  # fantasma
    s9 = wb.create_sheet("subform_9")
    s9.append(["Submission Id", "Sector", "Ubicación", "PK del defecto (Abscisa)", "Status OFF", "Status ON",
               "Forma N", "Forma S", "Forma E", "Forma O", "OL/RE", "Foto OL/RE", "Profundidad (M)",
               "Foto Profundidad", "Severidad (P/RE)", "Severidad (%IR)", "Clasificación ",
               "Forma del defecto", "Comentarios", "Caracter de la indicación"])
    s9.append(["sub-1", "ramal Marsella ", "4.959865,-75.7705421", "2+", None, None, None, None, None, None,
               None, None, None, None, None, None, "", "", None, ""])          # incompleto (se conserva)
    s9.append(["sub-2", None, None, None, None, None, None, None, None, None, None, None, None, None,
               None, None, "", "", None, ""])                                   # fantasma (fuera)
    s9.append(["sub-1", "ramal Marsella", "4.96,-75.78", "4+500", None, None, 10, 12, 8, 9, 35, "f", 1.2,
               "f", None, None, "", "", "defecto real", "Anódico-Anódico"])
    wb.save(path)
    return path


def _resist(path):
    wb = openpyxl.Workbook()
    _root(wb.active, "Juan Perez"); wb.active.title = "Root"
    s7 = wb.create_sheet("subform_7")
    s7.append(["Submission Id", "PK", "Sector", "Profundidad", "Ubicación", "Resistencia 1 metro",
               "Foto resistencia 1 metro", "Resistencia 2 metros", "Foto resistencia 2 metros",
               "Resistencia 3 metros", "Foto resistencia 3 metros"])
    s7.append(["sub-1", "7+000", "ramal Marsella ", "150", "4.9442265,-75.745302", "15", "f", "10", "f", "7", "f"])
    s7.append(["sub-2", None, None, None, None, None, None, None, None, None, None])       # fantasma
    wb.save(path)
    return path


def test_postes_con_pk_en_kilometros_y_sin_fila_fantasma(tmp_path):
    d = leer_dcvg_fastfield(_dcvg(tmp_path / "d.xlsx"))
    assert [p["pk_m"] for p in d["postes"]] == [3000, 7000]


def test_defecto_fantasma_fuera_e_incompleto_se_conserva(tmp_path):
    d = leer_dcvg_fastfield(_dcvg(tmp_path / "d.xlsx"))
    assert len(d["defectos"]) == 2
    assert d["defectos"][0]["pk_m"] is None                   # '2+' incompleto: lo avisa la app
    assert d["defectos"][1]["pk_m"] == 4500 and d["defectos"][1]["caracter"] == "AA"


def test_resistividad_fantasma_fuera(tmp_path):
    r = leer_resistividades_fastfield(_resist(tmp_path / "r.xlsx"))
    assert len(r) == 1 and r[0]["pk_m"] == 7000 and r[0]["r1"] == 15.0


def test_tecnico_que_es_una_foto_no_es_inspector(tmp_path):
    d = leer_dcvg_fastfield(_dcvg(tmp_path / "d.xlsx"))
    assert d["meta"]["tecnico"] == ""
    assert "inspector" not in info_desde_meta(d["meta"], [])


def test_tecnico_con_nombre_si_pasa(tmp_path):
    d = leer_dcvg_fastfield(_dcvg(tmp_path / "d.xlsx", tecnico="Juan Perez"))
    assert d["meta"]["tecnico"] == "Juan Perez"
    assert info_desde_meta(d["meta"], [])["inspector"] == "Juan Perez"
