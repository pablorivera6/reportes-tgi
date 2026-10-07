"""Gráfica Resistividad (DCVG): las líneas de criterio y el rango de los ejes
deben cubrir los datos reales.

La plantilla trae las líneas de criterio (Muy corrosivo 500, Corrosivo 1000,
Moderadamente 2000, Levemente 10000) ancladas en T10:T11 = 0..4000 (la
abscisa del ejemplo) y el eje Y logarítmico 100..1e7. Con un tramo en el
K 136+000 las líneas quedaban fuera, y `fill_graficas_dcvg` además recortaba
esas series como si fueran datos (el nombre de la hoja de datos 'Resistividad'
está contenido en 'Gráfica Resistividad'), dejándolas apuntando a una fila
vacía. El resultado: una gráfica con la leyenda de las capas (0-1, 1-2, 2-3 m)
y nada más."""
import os

import openpyxl
import pytest

from generator import ReportGenerator, resource_path

INFO = {'fecha': '12/03/2025', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Armenia',
        'tipo_ducto': 'Ramal', 'tipo_inspeccion': 'DCVG'}


def _resist():
    return [{"pk_m": 137000, "sector": "A", "profundidad": 1.2, "r1": 0.05, "r2": 0.04, "r3": 0.03},
            {"pk_m": 145000, "sector": "B", "profundidad": 1.0, "r1": 3.0, "r2": None, "r3": None},
            {"pk_m": 159000, "sector": "C", "profundidad": 1.2, "r1": 20.0, "r2": 15.0, "r3": 10.0}]


def _generar(tmp_path):
    gen = ReportGenerator(resource_path("DCVG_REP.xlsx"))
    postes = [{"tipo": "Poste de Potencial", "pk_m": pk, "on": -1600.0, "off": -1100.0}
              for pk in (136300, 150000, 160000)]
    resist = _resist()
    gen.fill_general_info(dict(INFO))
    gen.fill_dcvg(postes, [], resist)
    gen.fill_resistividad(resist)
    gen.fill_graficas_dcvg(gen.dcvg_filas, len(resist))
    out = os.path.join(tmp_path, "d.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)


def _series(ch):
    return [(s.xVal.numRef.f, s.yVal.numRef.f) for s in ch.series]


def test_las_series_de_datos_cubren_las_filas_escritas(tmp_path):
    wb = _generar(tmp_path)
    ch = wb["Gráfica Resistividad"]._charts[0]
    datos = [s for s in _series(ch) if s[0].startswith("Resistividad!")]
    assert len(datos) == 3
    for x, y in datos:
        assert x == "Resistividad!$A$9:$A$11", x
        assert y.endswith("$9:$" + y.split("$")[-2] + "$11"), y
    # las tres capas: ρC1, ρC2, ρC3 (columnas R, S, T)
    assert {y.split("!$")[1].split("$")[0] for _x, y in datos} == {"R", "S", "T"}


def test_las_lineas_de_criterio_no_se_recortan_como_datos(tmp_path):
    wb = _generar(tmp_path)
    ch = wb["Gráfica Resistividad"]._charts[0]
    crit = [s for s in _series(ch) if "Gráfica Resistividad" in s[0]]
    assert len(crit) == 4
    for x, _y in crit:
        assert x == "'Gráfica Resistividad'!$T$10:$T$11", x


def test_los_criterios_abarcan_el_recorrido_real(tmp_path):
    wb = _generar(tmp_path)
    ws = wb["Gráfica Resistividad"]
    ch = ws._charts[0]
    assert ws["T10"].value == ch.x_axis.scaling.min <= 137000
    assert ws["T11"].value == ch.x_axis.scaling.max >= 159000
    # los valores de los criterios siguen siendo los de la plantilla
    assert [ws.cell(10, c).value for c in range(21, 25)] == [500, 1000, 2000, 10000]


def test_el_eje_y_cubre_los_valores_de_resistividad(tmp_path):
    """ρ1 = 2π·0.05·100 ≈ 31 Ω·cm queda por debajo del mínimo 100 del
    template; ρ3 = 2π·10·300 ≈ 18 850 por encima de nada. El eje debe
    contenerlos (escala log, potencias de 10)."""
    wb = _generar(tmp_path)
    ch = wb["Gráfica Resistividad"]._charts[0]
    assert ch.y_axis.scaling.min is not None and ch.y_axis.scaling.min <= 31
    assert ch.y_axis.scaling.max is not None and ch.y_axis.scaling.max >= 18850
    assert ch.y_axis.scaling.logBase == 10


def test_resistividades_con_claves_de_la_api_tambien_se_escriben(tmp_path):
    """El adaptador de la API trae r_1m/r_2m/r_3m; el generador debe leerlas."""
    gen = ReportGenerator(resource_path("DCVG_REP.xlsx"))
    resist = [{"pk_m": 1000, "sector": "A", "profundidad": 1.2, "r_1m": 2.0, "r_2m": 1.5, "r_3m": 1.0}]
    gen.fill_resistividad(resist)
    out = os.path.join(tmp_path, "d.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)["Resistividad"]
    assert (ws["F9"].value, ws["H9"].value, ws["J9"].value) == (2.0, 1.5, 1.0)
