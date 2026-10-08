"""El informe DCVG lleva solo dos gráficas: la general de defectos (GRAFICA
DCVG) y la de resistividades (Gráfica Resistividad). Las hojas 'K xxx+xxx -
K yyy+yyy' con una gráfica cada ~5 km ya no se usan (decisión del ingeniero,
2026-10) y no deben salir, ni en la web ni en el escritorio."""
import io
import os

import openpyxl
from streamlit.testing.v1 import AppTest

from generator import ReportGenerator, resource_path

RAIZ = os.path.join(os.path.dirname(__file__), "..")
APP = os.path.join(RAIZ, "streamlit_app.py")


def _datos():
    # 136+300 .. 197+325: antes salían ~13 hojas de 5 km
    postes = [{"tipo": "Poste", "pk_m": pk, "on": -1600.0, "off": -1100.0,
               "lat": 4.9, "lon": -75.7} for pk in (136300, 140000, 145000,
               150000, 197325)]
    defectos = [{"pk_m": 142000, "forma_n": 1, "forma_s": 1, "forma_e": 1,
                 "forma_o": 1, "ol_re": 30, "profundidad": 190, "caracter": "CC",
                 "lat": 4.9, "lon": -75.7, "comentarios": "d"}]
    return postes, defectos


def test_web_genera_dcvg_sin_hojas_por_rango():
    at = AppTest.from_file(APP, default_timeout=180)
    at.run()
    postes, defectos = _datos()
    at.session_state.data["info"].update({"tipo_inspeccion": "DCVG",
                                          "tramo": "Ramal Armenia"})
    at.session_state.data["dcvg_postes"] = postes
    at.session_state.data["dcvg_defectos"] = defectos
    at.session_state["elaboro_nombre"] = "Pablo Rivera"
    at.run()
    boton = [b for b in at.button if str(b.label).strip().lower() == "generar informe"][0]
    boton.click().run()
    assert not at.exception, at.exception
    wb = openpyxl.load_workbook(io.BytesIO(at.session_state["informe_bytes"]))
    plantilla = openpyxl.load_workbook(resource_path("DCVG_REP.xlsx"))
    assert wb.sheetnames == plantilla.sheetnames
    con_grafica = [ws.title for ws in wb.worksheets if ws._charts]
    assert con_grafica == ["GRAFICA DCVG", "Gráfica Resistividad"]


def test_ningun_flujo_crea_hojas_por_rango():
    assert not hasattr(ReportGenerator, "fill_rangos_dcvg")
    for archivo in ("streamlit_app.py", "app.py"):
        src = open(os.path.join(RAIZ, archivo), encoding="utf-8").read()
        assert "fill_rangos_dcvg" not in src, archivo
