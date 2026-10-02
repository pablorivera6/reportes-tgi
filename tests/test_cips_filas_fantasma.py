"""Un export del logger puede traer una celda suelta al final de la hoja (un
espacio en la fila 1.048.576). pandas la lee como un millón de filas vacías;
al unir los archivos la tabla superaba el máximo de Excel y el proceso moría
con un error que no decía nada: 'At least one sheet must be visible'."""
import os

import openpyxl
import pandas as pd
import pytest

from mod_unificar import ejecutar_unificar


def _archivo(ruta, n, fantasma=None, desde=1):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey Data"
    ws.append(["Data No", "On Voltage", "Off Voltage", "Latitude", "Longitude"])
    for i in range(desde, desde + n):
        ws.append([i, -1.2, -0.9, 4.8, -74.9])
    if fantasma:
        ws.cell(fantasma, 1).value = " "
    d = wb.create_sheet("DCP Data")
    d.append(["Data No", "Device ID", "Comments"])
    d.append([desde, "Rectifier", "inicio"])
    if fantasma:
        d.cell(fantasma, 2).value = " "
    wb.save(ruta)


def test_filas_fantasma_no_tumban_la_unificacion(tmp_path):
    _archivo(os.path.join(tmp_path, "a.xlsx"), 10)
    _archivo(os.path.join(tmp_path, "b incre.xlsx"), 10, fantasma=1048576, desde=11)
    salida = ejecutar_unificar(str(tmp_path))
    survey = pd.read_excel(salida, sheet_name="Survey Data")
    assert len(survey) == 20
    assert sorted(survey["Data No"]) == list(range(1, 21))
    assert len(pd.read_excel(salida, sheet_name="DCP Data")) == 2


def test_fila_con_datos_reales_no_se_descarta(tmp_path):
    # Una fila con lectura pero sin Data No sigue siendo una lectura.
    ruta = os.path.join(tmp_path, "a.xlsx")
    _archivo(ruta, 3)
    wb = openpyxl.load_workbook(ruta)
    wb["Survey Data"].append([None, -1.1, -0.8, 4.8, -74.9])
    wb.save(ruta)
    survey = pd.read_excel(ejecutar_unificar(str(tmp_path)), sheet_name="Survey Data")
    assert len(survey) == 4


def test_demasiadas_filas_reales_da_un_error_claro(tmp_path, monkeypatch):
    import mod_unificar
    monkeypatch.setattr(mod_unificar, "MAX_FILAS_EXCEL", 15)
    _archivo(os.path.join(tmp_path, "a.xlsx"), 10)
    _archivo(os.path.join(tmp_path, "b.xlsx"), 10, desde=11)
    with pytest.raises(ValueError, match="filas"):
        ejecutar_unificar(str(tmp_path))
