"""Los 3 puntos de carga DCVG aceptan varios archivos y se combinan."""
import os

import openpyxl
import pytest

from dcvg_reader import (leer_dcvg_fastfield_varios,
                         leer_resistividades_fastfield_varios,
                         leer_hallazgos_logger_varios)

_REF = "/private/tmp/dcvg_ref"


def _ff(tmp_path, nombre, pk_defecto, tecnico=""):
    wb = openpyxl.Workbook()
    root = wb.active; root.title = "Root"
    root.append(["Contratista", "Fecha", "Cliente",
                 "Troncal o ramal inspeccionado", "Técnico a cargo"])
    root.append(["PCC", "07-28-2026", "TGI", "Montenegro", tecnico])
    s9 = wb.create_sheet("subform_9")
    s9.append(["Submission Id", "PK del defecto (Abscisa)", "Forma N", "Forma S",
               "Forma E", "Forma O", "OL/RE", "Caracter de la indicación",
               "Ubicación"])
    s9.append(["x", pk_defecto, 10, 10, 10, 10, 30, "Catódico-Catódico", "4.5,-75.7"])
    ruta = os.path.join(tmp_path, nombre)
    wb.save(ruta)
    return ruta


def test_dcvg_fastfield_varios_combina(tmp_path):
    a = _ff(tmp_path, "a.xlsx", "1+000", tecnico="")
    b = _ff(tmp_path, "b.xlsx", "2+000", tecnico="Juan Perez")
    d = leer_dcvg_fastfield_varios([a, b])
    pks = sorted(x["pk_m"] for x in d["defectos"])
    assert pks == [1000, 2000]
    # meta toma el técnico del archivo que lo trae
    assert d["meta"]["tecnico"] == "Juan Perez"


@pytest.mark.skipif(not os.path.isdir(_REF), reason="reales no disponibles")
def test_varios_montenegro_duplicado():
    f = os.path.join(_REF, "Dcvg Fastfield.xlsx")
    uno = leer_dcvg_fastfield_varios([f])
    dos = leer_dcvg_fastfield_varios([f, f])
    assert len(dos["defectos"]) == 2 * len(uno["defectos"])
    # las resistividades repetidas se descartan (mismo archivo dos veces)
    r = os.path.join(_REF, "Resistividades Fastfield.xlsx")
    assert len(leer_resistividades_fastfield_varios([r, r])) == \
        len(leer_resistividades_fastfield_varios([r]))
    lg = os.path.join(_REF, "DCVG MONTENEGRO 28 07 2026.xlsx")
    assert len(leer_hallazgos_logger_varios([lg, lg])) == \
        2 * len(leer_hallazgos_logger_varios([lg]))


def _resist(tmp_path, nombre, pks):
    wb = openpyxl.Workbook()
    s7 = wb.create_sheet("subform_7")
    s7.append(["Submission Id", "PK", "Sector", "Profundidad", "Ubicación",
               "Resistencia 1 metro", "Resistencia 2 metros",
               "Resistencia 3 metros"])
    for pk in pks:
        s7.append(["x", pk, "Potrero", 150, "4.9,-75.7", 1.5, 2.5, 3.5])
    ruta = os.path.join(tmp_path, nombre)
    wb.save(ruta)
    return ruta


def test_resistividades_varios_descarta_repetidas(tmp_path):
    """El mismo FastField cargado dos veces (o dos exportes solapados) no
    duplica los puntos: una resistividad idéntica entra una sola vez."""
    a = _resist(tmp_path, "a.xlsx", ["1+000", "2+000"])
    b = _resist(tmp_path, "b.xlsx", ["2+000", "3+000"])
    assert len(leer_resistividades_fastfield_varios([a, a])) == 2
    pks = sorted(x["pk_m"] for x in leer_resistividades_fastfield_varios([a, b]))
    assert pks == [1000, 2000, 3000]
