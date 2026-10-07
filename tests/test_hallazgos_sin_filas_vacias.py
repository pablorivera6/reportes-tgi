"""La hoja Hallazgos de las TRES plantillas trae el encabezado (ÍTEM, ABSCISA
INICIAL…) en la fila 11, pero el generador escribía desde la fila 18 quemada:
todo informe salía con seis filas vacías entre el encabezado y el primer
hallazgo (observación del ingeniero, Marsella 2026-10). La fila de inicio se
ubica por la etiqueta, nunca quemada."""
import os

import openpyxl
import pytest

from generator import ReportGenerator

PLANTILLAS = ["EN BLANCO.xlsx", "CIPS EN BLANCO.xlsx", "DCVG_REP.xlsx"]
INFO = {'fecha': '07/05/2026', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Marsella',
        'tipo_inspeccion': 'DCVG', 'contrato': '551007370', 'contratista': 'PCC',
        'ot': '1300', 'inspector': 'X'}
HALLAZGOS = [{'abscisa_val': 490, 'tipo': 'Válvula', 'descripcion': 'PK 0.490 válvula',
              'lat': 4.968, 'lon': -75.78, 'fecha': '07/05/2026'},
             {'abscisa_val': 90, 'tipo': 'Observación de campo', 'descripcion': 'PK 0.000 fin',
              'lat': 4.971, 'lon': -75.777, 'fecha': '07/05/2026'}]


def _hoja(tmp_path, plantilla, hallazgos):
    gen = ReportGenerator(plantilla)
    gen.fill_hallazgos(hallazgos, INFO)
    out = os.path.join(tmp_path, "h.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)['Hallazgos']


def _fila_encabezado(ws):
    return next(r for r in range(1, 40)
                if str(ws.cell(r, 1).value or '').strip().upper().startswith(('ÍTEM', 'ITEM')))


@pytest.mark.parametrize("plantilla", PLANTILLAS)
def test_el_primer_hallazgo_va_justo_debajo_del_encabezado(tmp_path, plantilla):
    ws = _hoja(tmp_path, plantilla, HALLAZGOS)
    enc = _fila_encabezado(ws)
    assert ws.cell(enc + 1, 1).value == 1
    assert ws.cell(enc + 1, 2).value == 90            # ordenado por abscisa
    assert ws.cell(enc + 2, 1).value == 2
    assert ws.cell(enc + 2, 2).value == 490
    assert ws.cell(enc + 2, 12).value == 'Válvula'
    # nada regado más abajo (la plantilla PAP traía un '1' de ejemplo en A18)
    for r in range(enc + 3, enc + 12):
        assert all(ws.cell(r, c).value in (None, '') for c in range(1, 14)), r


@pytest.mark.parametrize("plantilla", PLANTILLAS)
def test_sin_hallazgos_no_queda_residuo_de_la_plantilla(tmp_path, plantilla):
    ws = _hoja(tmp_path, plantilla, [])
    enc = _fila_encabezado(ws)
    for r in range(enc + 1, enc + 12):
        assert all(ws.cell(r, c).value in (None, '') for c in range(1, 14)), r
