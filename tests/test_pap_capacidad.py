"""La hoja 'Potenciales PAP' trae 63 filas de datos (12-74) y debajo el
bloque de firmas (fila 75-79, con celdas combinadas). Con más postes que eso
el generador usaba insert_rows, que en openpyxl NO corre las celdas
combinadas: los postes 64-68 caían dentro del bloque de firmas y se perdían,
y la numeración vieja de la plantilla quedaba regada debajo de los datos."""
import datetime as dt
import os

import openpyxl
import pytest

from generator import ReportGenerator


def _pots(n):
    return [{'abscisa': i * 1000, 'fecha': '09-27-2026',
             'ref_geografica': 'Poste de Potencial',
             'on_mv': -1000 - i, 'off_mv': -900 - i} for i in range(n)]


def _generar(tmp_path, n):
    gen = ReportGenerator()
    gen.fill_potenciales_pap(_pots(n))
    out = os.path.join(tmp_path, "pap.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)['Potenciales PAP']


def _fila_firmas(ws):
    for r in range(12, ws.max_row + 1):
        if ws.cell(r, 3).value == 'ELABORÓ':
            return r
    return None


@pytest.mark.parametrize("n", [10, 63, 69, 300])
def test_ningun_poste_se_pierde_y_las_firmas_quedan_debajo(tmp_path, n):
    ws = _generar(tmp_path, n)
    for i in range(n):
        r = 12 + i
        assert ws.cell(r, 1).value == i + 1
        assert ws.cell(r, 2).value == i * 1000, f"poste {i + 1} perdido"
        assert ws.cell(r, 5).value == -1000 - i
    firmas = _fila_firmas(ws)
    assert firmas is not None and firmas > 11 + n
    # el bloque conserva sus celdas combinadas y sus fórmulas
    rangos = {str(m) for m in ws.merged_cells.ranges}
    assert f"C{firmas}:K{firmas}" in rangos
    assert f"A{firmas - 1}:AA{firmas - 1}" in rangos
    assert ws.cell(firmas + 1, 3).value == '=Informe!D100'
    # ninguna celda combinada se mete en la zona de datos
    assert not any(m.min_row <= 11 + n and m.max_row >= 12
                   for m in ws.merged_cells.ranges)
    # nada suelto debajo del bloque de firmas
    for r in range(firmas + 4, ws.max_row + 1):
        assert all(ws.cell(r, c).value is None for c in range(1, 28)), r
    assert str(ws.print_area).endswith(f"$AA${firmas + 3}")


def test_filas_extra_con_formato_de_tabla(tmp_path):
    ws = _generar(tmp_path, 69)
    for r in (75, 80):
        assert ws.cell(r, 2).number_format == ws.cell(12, 2).number_format
        assert ws.cell(r, 2).border.bottom.style == 'thin'


def test_fecha_como_fecha_real(tmp_path):
    # FastField exporta '09-27-2026' (mes-día-año) como texto; la columna
    # está formateada como fecha, así que se escribe una fecha de verdad.
    ws = _generar(tmp_path, 3)
    v = ws.cell(12, 3).value
    assert isinstance(v, (dt.date, dt.datetime))
    assert (v.year, v.month, v.day) == (2026, 9, 27)


@pytest.mark.parametrize("texto,esperado", [
    ("09-27-2026", "27/09/2026"),   # FastField: mes-día-año
    ("09-05-2026", "05/09/2026"),   # ambiguo: se respeta el formato FastField
    ("27-09-2026", "27/09/2026"),   # ya venía día-mes-año
    ("", ""), ("sin fecha", "sin fecha"),
])
def test_lector_normaliza_fecha_fastfield(texto, esperado):
    from readers import _fecha_fastfield
    assert _fecha_fastfield(texto) == esperado


def test_item_no_se_parte_en_dos_lineas(tmp_path):
    ws = _generar(tmp_path, 69)
    assert not ws.cell(70, 1).alignment.wrap_text


def test_firmas_no_pisan_datos_de_postes(tmp_path):
    gen = ReportGenerator()
    gen.fill_potenciales_pap(_pots(69))
    p = {'nombre': 'ING X', 'cargo': 'Ingeniero', 'empresa': 'PCC'}
    gen.fill_firmas(p, p, p)
    out = os.path.join(tmp_path, "f.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)['Potenciales PAP']
    for r in range(12, 12 + 69):
        assert ws.cell(r, 4).value == 'Poste de Potencial', r
        assert ws.cell(r, 15).value is None and ws.cell(r, 24).value is None, r
