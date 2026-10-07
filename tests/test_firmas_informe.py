"""Las firmas (ELABORÓ / REVISÓ / APROBÓ) se escriben en el bloque de la hoja
Informe ubicado por etiqueta. fill_firmas escribía en las filas 104-106
quemadas, que en las tres plantillas están VACÍAS y debajo del bloque (PAP
100-102, CIPS 94-96, DCVG 99-101): el informe salía con las firmas de ejemplo
de la plantilla ('Alejandro Rivera', 'Protección Catódica de Colombia'). Las
demás hojas muestran las mismas firmas por fórmula hacia Informe."""
import os
import re

import openpyxl
import pytest

from generator import ReportGenerator

E = dict(nombre='ING ELABORA', cargo='Ingeniero Junior X', empresa='PCC Integrity')
R = dict(nombre='ING REVISA', cargo='Residente X', empresa='PCC Integrity')
A = dict(nombre='ING APRUEBA', cargo='Coordinador X', empresa='PCC Integrity')


def _bloque(ws):
    """{rol: (fila_nombre, col)} leyendo las etiquetas del bloque."""
    out = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper() in ('ELABORÓ', 'REVISÓ', 'APROBÓ'):
                col = c
                if str(ws.cell(r + 1, c).value or '').strip().lower().startswith(('nombre', 'cargo')):
                    col = next(cc for cc in range(c + 1, c + 8)
                               if type(ws.cell(r + 1, cc)).__name__ != 'MergedCell')
                out[v.strip().upper()] = (r + 1, col)
        if len(out) == 3:
            return out
    return out


@pytest.mark.parametrize("plantilla,hojas", [
    ("EN BLANCO.xlsx", ['Potenciales PAP', 'Hallazgos', 'Aislamientos']),
    ("DCVG_REP.xlsx", ['Inspección DCVG', 'Resistividad', 'Hallazgos']),
    ("CIPS EN BLANCO.xlsx", ['Potenciales PAP', 'Hallazgos', 'Aislamientos']),
])
def test_firmas_en_el_bloque_del_informe_y_por_formula_en_las_demas(tmp_path, plantilla, hojas):
    gen = ReportGenerator(plantilla)
    gen.fill_firmas(E, R, A)
    out = os.path.join(tmp_path, "f.xlsx")
    gen.save(out)
    wb = openpyxl.load_workbook(out)
    ws = wb['Informe']
    b = _bloque(ws)
    assert set(b) == {'ELABORÓ', 'REVISÓ', 'APROBÓ'}
    for rol, quien in (('ELABORÓ', E), ('REVISÓ', R), ('APROBÓ', A)):
        r, c = b[rol]
        assert [ws.cell(r + k, c).value for k in range(3)] == \
            [quien['nombre'], quien['cargo'], quien['empresa']], (plantilla, rol)
    # no queda ninguna firma de ejemplo de la plantilla
    textos = ' '.join(str(ws.cell(r, c).value) for r in range(ws.max_row - 20, ws.max_row + 1)
                      for c in range(1, ws.max_column + 1) if ws.cell(r, c).value)
    assert 'Alejandro Rivera' not in textos and 'Alejando Rivera' not in textos
    assert 'Protección Catódica de Col' not in textos
    # hojas secundarias: fórmula hacia la celda del Informe
    for hoja in hojas:
        wsh = wb[hoja]
        bh = _bloque(wsh)
        assert set(bh) == {'ELABORÓ', 'REVISÓ', 'APROBÓ'}, hoja
        for rol in bh:
            (r, c), (ri, ci) = bh[rol], b[rol]
            for k in range(3):
                f = str(wsh.cell(r + k, c).value)
                m = re.match(r"^='?Informe'?!\$?([A-Z]+)\$?(\d+)$", f)
                assert m, (hoja, rol, k, f)
                assert ws[f"{m.group(1)}{m.group(2)}"].coordinate == ws.cell(ri + k, ci).coordinate, (hoja, rol, k)


def test_las_dos_apps_firman_tambien_el_dcvg():
    """El flujo DCVG (web y escritorio) no llamaba a fill_firmas."""
    for archivo, inicio, fin in (("streamlit_app.py", "Rama DCVG", "tipo_inspeccion') == 'CIPS'"),
                                 ("app.py", "Rama DCVG", "Guardando informe DCVG")):
        src = open(archivo, encoding='utf-8').read()
        rama = src[src.index(inicio):src.index(fin, src.index(inicio))]
        assert 'fill_firmas(' in rama, archivo
