"""Ciclo de interrupción del informe (decisión del ingeniero, 2026-10-08):
PAP y CIPS se miden siempre con ON 1,6 s / OFF 0,4 s; en DCVG el ciclo está
pendiente de definir y la casilla va vacía. Lo fija el generador por tipo, así
que vale igual en la web y en el escritorio, y lo que trajera el equipo
('1600/400 ms') o la sesión no lo cambia."""
import os

import openpyxl
import pytest

import datos_tramo
from generator import ReportGenerator


def _valor_ciclo(ws):
    for r in range(6, 10):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower().startswith('ciclo'):
                for cc in range(c + 1, c + 8):
                    if type(ws.cell(r, cc)).__name__ != 'MergedCell':
                        return ws.cell(r, cc).value
    raise AssertionError("no se encontró la etiqueta Ciclo")


@pytest.mark.parametrize("plantilla,tipo,esperado", [
    ("EN BLANCO.xlsx", "PAP", "ON 1,6 s / OFF 0,4 s"),
    ("CIPS EN BLANCO.xlsx", "CIPS", "ON 1,6 s / OFF 0,4 s"),
    ("DCVG_REP.xlsx", "DCVG", None),
])
def test_ciclo_por_tipo_en_el_informe(tmp_path, plantilla, tipo, esperado):
    gen = ReportGenerator(plantilla)
    gen.fill_general_info({'tipo_inspeccion': tipo, 'tramo': 'Ramal Marsella',
                           'ciclo': '1600/400 ms'})          # lo del equipo no manda
    out = os.path.join(tmp_path, "c.xlsx")
    gen.save(out)
    v = _valor_ciclo(openpyxl.load_workbook(out)['Informe'])
    assert (v or None) == esperado


def test_ciclo_de():
    assert datos_tramo.ciclo_de("PAP") == "ON 1,6 s / OFF 0,4 s"
    assert datos_tramo.ciclo_de("cips") == "ON 1,6 s / OFF 0,4 s"
    assert datos_tramo.ciclo_de("DCVG") == ""
    assert datos_tramo.ciclo_pendiente("DCVG") and not datos_tramo.ciclo_pendiente("PAP")


@pytest.mark.parametrize("tipo,esperado", [("PAP", "ON 1,6 s / OFF 0,4 s"),
                                           ("CIPS", "ON 1,6 s / OFF 0,4 s"),
                                           ("DCVG", "")])
def test_autollenar_trae_el_ciclo_del_tipo(tipo, esperado):
    assert datos_tramo.autollenar("Ramal Marsella", tipo).get("ciclo") == esperado
