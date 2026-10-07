"""La plantilla DCVG trae quemados en la hoja Informe el punto inicial (0), el
punto final (4000) y la longitud (4 km) del informe de ejemplo, y la gráfica
DCVG con el eje X de 0 a 4000. El generador nunca los escribía: TODOS los
informes DCVG salían con "longitud inspeccionada 4 km" y la gráfica cortada
en el K 004+000, aunque el tramo tuviera 60 km."""
import os

import openpyxl

from generator import ReportGenerator, resource_path

INFO = {'fecha': '12/03/2025', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Armenia',
        'tipo_ducto': 'Ramal', 'contrato': '551007370', 'ot': '1300013506',
        'tipo_inspeccion': 'DCVG'}


def _datos():
    postes = [{"tipo": "Poste de Potencial", "pk_m": pk, "on": -1600.0, "off": -1100.0,
               "lat": 4.90 + i / 100, "lon": -75.70 - i / 100}
              for i, pk in enumerate((136300, 140000, 145000, 150000, 160000))]
    defectos = [{"pk_m": 142000, "forma_n": 1, "forma_s": 1, "forma_e": 1, "forma_o": 1,
                 "ol_re": 30, "profundidad": 190, "caracter": "CC", "lat": 4.91, "lon": -75.71},
                {"pk_m": 161500, "forma_n": 1, "forma_s": 1, "forma_e": 1, "forma_o": 1,
                 "ol_re": 80, "profundidad": 150, "caracter": "AA", "lat": 4.95, "lon": -75.75}]
    resist = [{"pk_m": 137000, "sector": "A", "profundidad": 1.2, "r1": 1.0, "r2": 0.8, "r3": 0.6},
              {"pk_m": 159000, "sector": "B", "profundidad": 1.2, "r1": 2.0, "r2": 1.5, "r3": 1.0}]
    return postes, defectos, resist


def _generar(tmp_path, info=INFO):
    gen = ReportGenerator(resource_path("DCVG_REP.xlsx"))
    postes, defectos, resist = _datos()
    gen.fill_general_info(dict(info))
    gen.fill_dcvg(postes, defectos, resist)
    gen.fill_resistividad(resist)
    gen.fill_graficas_dcvg(gen.dcvg_filas, len(resist))
    out = os.path.join(tmp_path, "d.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)


def _celda_tras_etiqueta(ws, etiqueta, col):
    for r in range(25, 45):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and etiqueta in v.upper():
            # el valor está en la misma fila o en la siguiente (la etiqueta
            # 'PUNTO INICIAL' ocupa A30:F31 y su valor es G31)
            for rr in (r, r + 1):
                val = ws.cell(rr, col).value
                if val not in (None, ''):
                    return val
            return None
    raise AssertionError(f"no está la etiqueta {etiqueta}")


def test_punto_inicial_final_y_longitud_son_los_del_recorrido(tmp_path):
    ws = _generar(tmp_path)["Informe"]
    assert _celda_tras_etiqueta(ws, "PUNTO INICIAL", 7) == 136300
    assert _celda_tras_etiqueta(ws, "PUNTO FINAL", 7) == 161500
    # el valor va en la celda combinada G31 (la etiqueta ocupa A30:F31), no en G30
    assert ws["G31"].value == 136300 and ws["G30"].value in (None, "")
    # longitud enterrada total e inspeccionada = extensión real (25.2 km), no 4
    total = _celda_tras_etiqueta(ws, "LONGITUD TOTAL DE LA LÍNEA ENTERRADA", 7)
    insp = _celda_tras_etiqueta(ws, "LONGITUD TOTAL DE LA LÍNEA ENTERRADA", 16)
    assert abs(float(total) - 25.2) < 0.01 and abs(float(insp) - 25.2) < 0.01


def test_longitud_total_de_datos_generales_manda_si_viene(tmp_path):
    ws = _generar(tmp_path, dict(INFO, longitud_km=45.334))["Informe"]
    total = _celda_tras_etiqueta(ws, "LONGITUD TOTAL DE LA LÍNEA ENTERRADA", 7)
    insp = _celda_tras_etiqueta(ws, "LONGITUD TOTAL DE LA LÍNEA ENTERRADA", 16)
    assert abs(float(total) - 45.334) < 0.001        # la del tramo (consolidado)
    assert abs(float(insp) - 25.2) < 0.01            # lo recorrido


def test_coordenadas_de_los_extremos_y_sin_altura_de_ejemplo(tmp_path):
    ws = _generar(tmp_path)["Informe"]
    # lat/lon del primer y último punto con coordenadas (postes/defectos)
    assert abs(float(ws["P31"].value) - 4.90) < 1e-6 and abs(float(ws["V31"].value) + 75.70) < 1e-6
    assert abs(float(ws["P32"].value) - 4.95) < 1e-6 and abs(float(ws["V32"].value) + 75.75) < 1e-6
    # la altura del ejemplo (2443 msnm / 2053 msnm) no se queda
    for c in ("AB31", "AB32"):
        assert "msnm" not in str(ws[c].value or "")


def test_graficas_dcvg_cubren_todo_el_recorrido(tmp_path):
    wb = _generar(tmp_path)
    ch = wb["GRAFICA DCVG"]._charts[0]
    assert ch.x_axis.scaling.min is not None and ch.x_axis.scaling.min <= 136300
    assert ch.x_axis.scaling.max is not None and ch.x_axis.scaling.max >= 161500
    assert ch.x_axis.scaling.max < 200000             # no se queda en el 4000 ni se dispara
    chr_ = wb["Gráfica Resistividad"]._charts[0]
    assert chr_.x_axis.scaling.min <= 137000 and chr_.x_axis.scaling.max >= 159000


def test_sin_fill_general_info_tampoco_queda_el_4(tmp_path):
    """El flujo puede llamar fill_dcvg sin info: aun así no deben quedar los
    valores del ejemplo."""
    gen = ReportGenerator(resource_path("DCVG_REP.xlsx"))
    postes, defectos, resist = _datos()
    gen.fill_dcvg(postes, defectos, resist)
    out = os.path.join(tmp_path, "d.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)["Informe"]
    assert _celda_tras_etiqueta(ws, "PUNTO FINAL", 7) == 161500
    assert float(_celda_tras_etiqueta(ws, "LONGITUD TOTAL DE LA LÍNEA ENTERRADA", 16)) > 20
