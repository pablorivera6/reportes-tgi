"""Hoja Aislamientos: había DOS fill_aislamientos en generator.py; la que
quedaba activa insertaba filas (insert_rows) y, como en Potenciales PAP,
openpyxl no corre las celdas combinadas del bloque de firmas: con más de una
junta el bloque quedaba encima de los datos. Además las dos usaban nombres de
campo distintos. Ahora hay una sola, con la capacidad de la plantilla (5 filas
en PAP, 19 en CIPS) y bajando el bloque si no caben."""
import os

import openpyxl
import pytest

from generator import ReportGenerator


def _juntas(n):
    return [dict(abscisado=1000 * (i + 1), tag=f'JA-{i + 1}', clase='150', diametro='4',
                 presion='-', temperatura='-', tipo_brida='WN', numero_pernos=8,
                 diametro_pernos='5/8', tipo_aislamiento='Kit', porcentaje_aislamiento=99,
                 pot_on_arriba=-1500, pot_off_arriba=-1000, pot_on_abajo=-700,
                 pot_off_abajo=-650, diferencia=800, diagnostico='aislado',
                 latitud=4.96 + i * 0.001, longitud=-75.77, observaciones='ok') for i in range(n)]


def _hoja(tmp_path, plantilla, n):
    gen = ReportGenerator(plantilla)
    gen.fill_aislamientos(_juntas(n))
    gen.fill_firmas({'nombre': 'E'}, {'nombre': 'R'}, {'nombre': 'A'})
    out = os.path.join(tmp_path, "a.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)['Aislamientos']


def _fila_firmas(ws):
    return next(r for r in range(13, ws.max_row + 1) if ws.cell(r, 3).value == 'ELABORÓ')


@pytest.mark.parametrize("plantilla,n", [("EN BLANCO.xlsx", 1), ("EN BLANCO.xlsx", 5),
                                         ("EN BLANCO.xlsx", 9), ("CIPS EN BLANCO.xlsx", 25)])
def test_todas_las_juntas_y_el_bloque_de_firmas_debajo(tmp_path, plantilla, n):
    ws = _hoja(tmp_path, plantilla, n)
    for i in range(n):
        r = 13 + i
        assert ws.cell(r, 1).value == i + 1
        assert ws.cell(r, 2).value == 1000 * (i + 1), f"junta {i + 1} perdida"
        assert ws.cell(r, 3).value == f'JA-{i + 1}'
        assert ws.cell(r, 9).value == 8 and ws.cell(r, 10).value == '5/8'
        assert ws.cell(r, 12).value == 99 and ws.cell(r, 17).value == 800
        assert ws.cell(r, 20).value == pytest.approx(4.96 + i * 0.001) and ws.cell(r, 21).value == -75.77
        assert ws.cell(r, 19).value == 'Aislado' and ws.cell(r, 22).value == 'Ok'
    firmas = _fila_firmas(ws)
    assert firmas > 12 + n
    rangos = {str(m) for m in ws.merged_cells.ranges}
    assert f"C{firmas}:G{firmas}" in rangos and f"H{firmas + 1}:Q{firmas + 1}" in rangos
    assert not any(m.min_row <= 12 + n and m.max_row >= 13 for m in ws.merged_cells.ranges)
    # firmas por fórmula hacia Informe, y nada regado entre datos y bloque
    assert str(ws.cell(firmas + 1, 3).value).startswith('=Informe!')
    for r in range(13 + n, firmas):
        assert all(ws.cell(r, c).value in (None, '') for c in range(1, 23)), r


def test_acepta_tambien_los_nombres_de_campo_del_adaptador(tmp_path):
    gen = ReportGenerator()
    gen.fill_aislamientos([dict(abscisa_val=500, tag='J', num_pernos=4, diam_pernos='1/2',
                                pct_aislamiento=98, dif_on=10, dif_off=5, lat=4.9, lon=-75.7)])
    out = os.path.join(tmp_path, "b.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)['Aislamientos']
    assert [ws.cell(13, c).value for c in (2, 9, 10, 12, 17, 18, 20, 21)] == \
        [500, 4, '1/2', 98, 10, 5, 4.9, -75.7]


def test_una_sola_definicion():
    src = open('generator.py', encoding='utf-8').read()
    assert src.count('def fill_aislamientos(') == 1
