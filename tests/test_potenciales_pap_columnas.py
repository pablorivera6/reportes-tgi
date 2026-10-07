"""La hoja 'Potenciales PAP' NO tiene las mismas columnas en las dos plantillas:
la de PAP trae 'Altura' en V y la de CIPS no, así que pintura/conexiones/
verticalidad/mantenimiento/observaciones están corridas una columna entre
ambas. El generador escribía columnas quemadas (las de PAP) y en un informe
CIPS los postes salían con la altura bajo ESTADO PINTURA y las observaciones
fuera de la tabla (observación del ingeniero: 'columnas corridas en los
postes'). Cada valor se ubica por el encabezado de SU plantilla."""
import os

import openpyxl
import pytest

from generator import ReportGenerator

POSTE = dict(abscisa=1000, fecha='05/07/2026', ref_geografica='Poste de Potencial',
             on_mv=-1500, off_mv=-1000, on_mv_neg2=-1400, off_mv_neg2=-900,
             on_mv_foraneo1=-10, off_mv_foraneo1=-11, on_mv_foraneo2=-20, off_mv_foraneo2=-21,
             potencial_natural=-600, polarizacion=-400, vac=1.5, resistencia=0.3, ir_on_off=500,
             lat=4.95, lon=-75.77, alt=1200, pintura='Bueno', conexiones='Regular',
             verticalidad='Malo', tipo_mant='Pintura', observaciones='obs poste')

ESPERADO = {  # (etiqueta fila 10, subetiqueta fila 11) -> valor
    ('ABSCISADO', None): 1000, ('REFERENCIA GEOGRÁFICA', None): 'Poste de Potencial',
    ('POTENCIAL NEGATIVO 1 TGI', 'ON'): -1500, ('POTENCIAL NEGATIVO 1 TGI', 'OFF'): -1000,
    ('POTENCIAL NEGATIVO 2 TGI', 'ON'): -1400, ('POTENCIAL NEGATIVO 2 TGI', 'OFF'): -900,
    ('POTENCIAL NEGATIVO 1 FORÁNEO', 'ON'): -10, ('POTENCIAL NEGATIVO 1 FORÁNEO', 'OFF'): -11,
    ('POTENCIAL NEGATIVO 2 FORÁNEO', 'ON'): -20, ('POTENCIAL NEGATIVO 2 FORÁNEO', 'OFF'): -21,
    ('POTENCIAL NATURAL', None): -600, ('POLARIZACIÓN', None): -400, ('VOLTAJE AC', None): 1.5,
    ('RESISTENCIA', None): 0.3, ('IR ON-OFF', None): 500, ('LATITUD', None): 4.95,
    ('LONGITUD', None): -75.77, ('ESTADO PINTURA', None): 'Bueno',
    ('ESTADO CONEXIONES', None): 'Regular', ('ESTADO VERTICALIDAD', None): 'Malo',
    ('TIPO DE MANTENIMIENTO', None): 'Pintura', ('OBSERVACIONES', None): 'Obs poste',
}


def _columnas(ws):
    """{(etiqueta, sub): col} leído del encabezado real (filas 10/11)."""
    out, actual = {}, None
    for c in range(1, ws.max_column + 1):
        lab = ws.cell(10, c).value
        if lab:
            actual = str(lab).replace('\n', ' ').strip().upper()
        sub = str(ws.cell(11, c).value or '').upper()
        if actual is None:
            continue
        clave = 'OFF' if 'OFF' in sub else ('ON' if 'ON' in sub else None)
        out.setdefault((actual, clave), c)   # la primera columna de cada etiqueta
    return out


@pytest.mark.parametrize("plantilla", ["EN BLANCO.xlsx", "CIPS EN BLANCO.xlsx"])
def test_cada_valor_bajo_su_encabezado(tmp_path, plantilla):
    gen = ReportGenerator(plantilla)
    gen.fill_potenciales_pap([POSTE], '05/07/2026')
    out = os.path.join(tmp_path, "p.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)['Potenciales PAP']
    cols = _columnas(ws)
    for (etq, sub), valor in ESPERADO.items():
        col = next((c for (e, s), c in cols.items() if e.startswith(etq) and s == sub
                    and not (etq == 'POTENCIAL NEGATIVO 1 TGI' and 'CORREGIDO' in e)), None)
        assert col is not None, (plantilla, etq, sub)
        assert ws.cell(12, col).value == valor, (plantilla, etq, sub)
    # la altura solo existe en la plantilla PAP; en CIPS no se escribe en ningún lado
    col_alt = next((c for (e, s), c in cols.items() if e.startswith('ALTURA')), None)
    if col_alt:
        assert ws.cell(12, col_alt).value == 1200
    ultima = max(cols.values())
    assert all(ws.cell(12, c).value in (None, '') for c in range(ultima + 1, ultima + 3)), \
        "nada escrito a la derecha de la tabla"
