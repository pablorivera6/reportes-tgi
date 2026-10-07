"""DOCUMENTOS DE REFERENCIA (hoja Informe): la TM0497 unificada a la referencia
controlada disponible y sin repetirse.

Observación del revisor (2026-10): el informe citaba NACE TM0497-2018-SG y la
referencia disponible es AMPP TM0497-2022; en la plantilla DCVG además salía
dos veces. Solo se corrige esa norma; las demás quedan como las trae cada
plantilla."""
import os

import openpyxl
import pytest

from generator import ReportGenerator, resource_path

INFO = {'fecha': '12/03/2025', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Salento',
        'tipo_ducto': 'Ramal', 'contrato': '551007370', 'ot': '1300012786'}


def _refs(tmp_path, plantilla, tipo):
    gen = ReportGenerator(resource_path(plantilla))
    gen.fill_general_info(dict(INFO, tipo_inspeccion=tipo))
    out = os.path.join(tmp_path, "i.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)["Informe"]
    ini = next(r for r in range(1, 40) if str(ws.cell(r, 1).value or '').strip().upper() == 'DOCUMENTOS DE REFERENCIA')
    fin = next(r for r in range(ini + 1, 40) if str(ws.cell(r, 1).value or '').strip().upper() == 'EQUIPOS UTILIZADOS')
    return [str(ws.cell(r, 1).value) for r in range(ini + 1, fin) if ws.cell(r, 1).value not in (None, '')]


@pytest.mark.parametrize("plantilla,tipo", [("EN BLANCO.xlsx", "PAP"),
                                            ("CIPS EN BLANCO.xlsx", "CIPS"),
                                            ("DCVG_REP.xlsx", "DCVG")])
def test_tm0497_unificada_y_sin_repetir(tmp_path, plantilla, tipo):
    refs = _refs(tmp_path, plantilla, tipo)
    texto = "\n".join(refs)
    assert sum('TM0497' in r for r in refs) == 1, refs
    assert "AMPP TM0497-2022" in texto and "2018" not in texto
    assert "SP0169-2024" in texto                       # lo demás de la plantilla sigue
    assert f"INSPECCIÓN {tipo} DE SPC" in texto         # PR-I-06 con el tipo real
    assert len(refs) == len(set(refs))                  # ninguna repetida


def test_pap_y_cips_conservan_la_sp0207(tmp_path):
    for plantilla, tipo in (("EN BLANCO.xlsx", "PAP"), ("CIPS EN BLANCO.xlsx", "CIPS")):
        assert any("SP0207-2007" in r for r in _refs(tmp_path, plantilla, tipo))


def test_dcvg_queda_con_tres_lineas_sin_huecos(tmp_path):
    refs = _refs(tmp_path, "DCVG_REP.xlsx", "DCVG")
    assert len(refs) == 3 and refs[-1].startswith("PR-I-06")
