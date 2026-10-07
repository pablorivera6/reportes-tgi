"""ANTECEDENTES y HUELLA OSCILOSCÓPICA se ocultan en la hoja Informe (pedido del
ingeniero, 2026-10): van vacías en los informes de PCC y ocupaban media
primera hoja. Se ocultan las filas (Excel no las imprime); el resto de la hoja
no se toca."""
import os

import openpyxl
import pytest

from generator import ReportGenerator, resource_path

INFO = {'fecha': '12/03/2025', 'gasoducto': 'Mariquita-Cali', 'tramo': 'Ramal Salento',
        'tipo_ducto': 'Ramal', 'contrato': '551007370', 'ot': '1300012786'}


def _informe(tmp_path, plantilla, tipo):
    gen = ReportGenerator(resource_path(plantilla))
    gen.fill_general_info(dict(INFO, tipo_inspeccion=tipo))
    out = os.path.join(tmp_path, "i.xlsx")
    gen.save(out)
    return openpyxl.load_workbook(out)["Informe"]


def _ocultas(ws):
    return {r for r, d in ws.row_dimensions.items() if d.hidden}


def _fila(ws, texto, col=1):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip().upper().startswith(texto):
            return r
    raise AssertionError(texto)


@pytest.mark.parametrize("plantilla,tipo", [("EN BLANCO.xlsx", "PAP"),
                                            ("CIPS EN BLANCO.xlsx", "CIPS"),
                                            ("DCVG_REP.xlsx", "DCVG")])
def test_antecedentes_oculto_y_lo_demas_visible(tmp_path, plantilla, tipo):
    ws = _informe(tmp_path, plantilla, tipo)
    oc = _ocultas(ws)
    r = _fila(ws, "ANTECEDENTES")
    assert r in oc and (r + 1) in oc
    # lo que sigue (SISTEMA INSPECCIONADO / PUNTO INICIAL) sigue visible
    for et in ("DESCRIPCIÓN DE LA LÍNEA", "PUNTO INICIAL", "CONCLUSIONES", "RECOMENDACIONES", "OBJETIVO"):
        assert _fila(ws, et) not in oc, et


@pytest.mark.parametrize("plantilla,tipo", [("EN BLANCO.xlsx", "PAP"),
                                            ("CIPS EN BLANCO.xlsx", "CIPS")])
def test_huella_oscilosopica_oculta_hasta_comentarios(tmp_path, plantilla, tipo):
    ws = _informe(tmp_path, plantilla, tipo)
    oc = _ocultas(ws)
    r_h = _fila(ws, "HUELLA OSCILOSC")
    r_c = _fila(ws, "COMENTARIOS")
    r_p = _fila(ws, "PARÁMETROS OPERATIVOS", col=2)
    assert all(r in oc for r in range(r_h, r_c + 1)), "huella y comentario ocultos"
    assert r_p not in oc and (r_p - 1) not in oc          # la tabla de URPC sigue visible
    # MONITOREO DE POTENCIALES (justo antes) sigue visible
    assert _fila(ws, "MONITOREO DE POTENCIALES") not in oc
    assert _fila(ws, "DATOS/KM") not in oc


def test_dcvg_no_tiene_huella_y_no_oculta_de_mas(tmp_path):
    ws = _informe(tmp_path, "DCVG_REP.xlsx", "DCVG")
    oc = _ocultas(ws)
    r = _fila(ws, "ANTECEDENTES")
    nuevas = {x for x in oc if x < 60}            # la plantilla ya trae ocultas 78-97
    assert nuevas == {r, r + 1}, nuevas


def test_se_puede_desactivar(tmp_path):
    gen = ReportGenerator(resource_path("EN BLANCO.xlsx"))
    gen.OCULTAR_SECCIONES = ()
    gen.fill_general_info(dict(INFO, tipo_inspeccion='PAP'))
    out = os.path.join(tmp_path, "i.xlsx")
    gen.save(out)
    ws = openpyxl.load_workbook(out)["Informe"]
    assert _fila(ws, "ANTECEDENTES") not in _ocultas(ws)
