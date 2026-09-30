"""Lee un informe CIPS **ya generado** (.xlsx) y lo devuelve en la forma que
produce el generador, para poder publicarlo al portal con
`db.guardar_inspeccion_cips`. Es el par CIPS de `informe_dcvg`.

La hoja 'Potenciales CIPS' se lee **por etiqueta de encabezado** (fila de
'ABSCISADO' + sub-fila ON/OFF), no por posición fija.

Las columnas E/F del informe son el potencial **oficial** (ya suavizado por
`cips_lrs._suavizar_outliers`); el crudo no viaja en el Excel. Por eso se
cargan como `on_limpio`/`off_limpio` —lo que usa el portal— y también como
`on_mv`/`off_mv`, para que el portal muestre exactamente lo del documento
entregado.
"""
import os

import openpyxl

from historicos import _num, _txt
from informe_dcvg import _fila_con, _hallazgos, _info as _info_dcvg

HOJA_CIPS = 'Potenciales CIPS'
HOJA_TRAMOS = 'Inv. Tramos no Inpeccionados'

# (etiqueta del encabezado, sub-etiqueta ON/OFF o None) → clave del generador
_COLUMNAS = {
    ('abscisado', None): 'abscisa_val',
    ('fecha', None): 'fecha',
    ('referencia geografica', None): 'referencia',
    ('potencial negativo 1 tgi', 'on'): 'on_limpio',
    ('potencial negativo 1 tgi', 'off'): 'off_limpio',
    ('potencial natural', None): 'natural_mv',
    ('polarizacion', None): 'polarizacion_mv',
    ('voltaje ac', None): 'vac',
    ('metal ir', 'on'): 'metal_on',
    ('metal ir', 'off'): 'metal_off',
    ('potencial lejano', 'on'): 'far_on',
    ('potencial lejano', 'off'): 'far_off',
    ('potencial cercano', 'on'): 'near_on',
    ('potencial cercano', 'off'): 'near_off',
    ('ir on-off', None): 'ir_on_off',
    ('latitud', None): 'lat',
    ('longitud', None): 'lon',
    ('observaciones', None): 'observaciones',
}
_TEXTO = {'fecha', 'referencia', 'observaciones'}


def _mapa(ws, enc):
    """{índice de columna (0-based): clave} leyendo la fila `enc` y la de
    abajo. Una etiqueta combinada (ON/OFF) manda hasta la siguiente."""
    mapa, arriba = {}, ''
    for c in range(1, ws.max_column + 1):
        v = _txt(ws.cell(row=enc, column=c).value)
        if v:
            arriba = v
        sub = _txt(ws.cell(row=enc + 1, column=c).value)
        sub = 'on' if sub.startswith('on') else 'off' if sub.startswith('off') else None
        if not v and sub is None:
            continue       # columna sin etiqueta propia ni ON/OFF: no arrastrar
        for (etq, s), clave in _COLUMNAS.items():
            # 'corregido' tiene su propia columna vacía: no confundirla
            if (arriba.startswith(etq) and 'corregido' not in arriba
                    and s == sub and clave not in mapa.values()):
                mapa[c - 1] = clave
                break
    return mapa


def _fecha_txt(v):
    if hasattr(v, 'strftime'):
        return v.strftime('%d/%m/%Y')
    return str(v).strip() if v not in (None, '') else ''


def _lecturas(wb):
    ws = wb[HOJA_CIPS]
    enc = _fila_con(ws, 'abscisado', col=2)
    if enc is None:
        raise ValueError(f"'{HOJA_CIPS}' sin encabezado ABSCISADO")
    mapa = _mapa(ws, enc)
    if 'abscisa_val' not in mapa.values() or 'off_limpio' not in mapa.values():
        raise ValueError(f"'{HOJA_CIPS}': no se encontraron ABSCISADO y OFF")
    out = []
    for fila in ws.iter_rows(min_row=enc + 2, values_only=True):
        p = {}
        for i, clave in mapa.items():
            v = fila[i] if i < len(fila) else None
            if clave == 'fecha':
                p[clave] = _fecha_txt(v)
            elif clave in _TEXTO:
                p[clave] = str(v).strip() if v not in (None, '') else ''
            else:
                p[clave] = _num(v)
        if p.get('abscisa_val') is None:
            # fin de datos: fila vacía o bloque de firmas
            if p.get('off_limpio') is None and p.get('on_limpio') is None:
                break
            continue
        p['abscisa_val'] = int(round(p['abscisa_val']))
        p['on_mv'], p['off_mv'] = p.get('on_limpio'), p.get('off_limpio')
        p['vac_mv'] = p.get('vac')
        out.append(p)
    return out


def _tramos_no_inspeccionados(wb):
    """Mismo formato que Hallazgos, con JUSTIFICACIÓN en la columna L."""
    if HOJA_TRAMOS not in wb.sheetnames:
        return []
    ws = wb[HOJA_TRAMOS]
    enc = _fila_con(ws, 'item', hasta=25)
    if enc is None:
        return []
    out = []
    for fila in ws.iter_rows(min_row=enc + 1, values_only=True):
        if any(_txt(x).startswith(('elaboro', 'reviso', 'aprobo'))
               for x in fila if isinstance(x, str)):
            break
        def v(i):
            return fila[i] if i < len(fila) else None
        ini = _num(v(1))
        just = str(v(11) or '').strip()
        if ini is None and not just:
            continue
        out.append({'abscisa_inicio': ini, 'abscisa_fin': _num(v(2)),
                    'longitud': _num(v(3)),
                    'lat_inicio': _num(v(6)), 'lon_inicio': _num(v(7)),
                    'lat_fin': _num(v(8)), 'lon_fin': _num(v(9)),
                    'fecha': v(10), 'justificacion': just})
    return out


def leer_informe_cips(ruta):
    """{info, cips, hallazgos, tramos, fuente}."""
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=False)
    try:
        if HOJA_CIPS not in wb.sheetnames:
            raise ValueError(f"{os.path.basename(ruta)}: no tiene la hoja "
                             f"'{HOJA_CIPS}' (¿es un informe CIPS?)")
        info = _info_dcvg(wb)
        info['tipo_inspeccion'] = 'CIPS'
        return {'info': info, 'cips': _lecturas(wb),
                'hallazgos': _hallazgos(wb),
                'tramos': _tramos_no_inspeccionados(wb),
                'fuente': os.path.basename(ruta)}
    finally:
        wb.close()
