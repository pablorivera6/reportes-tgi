"""Lee un informe PAP **ya generado** (.xlsx) y lo devuelve en la forma que
produce el generador, para poder publicarlo al portal con
`db.guardar_inspeccion_pap`. Es el par PAP de `informe_dcvg`/`informe_cips`.

La hoja 'Potenciales PAP' se lee **por etiqueta de encabezado**. Un poste sin
lectura (p.ej. "poste con avispas") se conserva: tiene abscisa y observación
aunque no traiga potenciales.
"""
import os

import openpyxl

from historicos import _num
from informe_cips import _fecha_txt, _mapa
from informe_dcvg import _fila_con, _hallazgos, _info as _info_dcvg

HOJA_PAP = 'Potenciales PAP'

# (etiqueta del encabezado, sub-etiqueta ON/OFF o None) → clave del generador
_COLUMNAS = {
    ('abscisado', None): 'abscisa',
    ('fecha', None): 'fecha',
    ('referencia geografica', None): 'ref_geografica',
    ('potencial negativo 1 tgi', 'on'): 'on_mv',
    ('potencial negativo 1 tgi', 'off'): 'off_mv',
    ('potencial natural', None): 'potencial_natural',
    ('polarizacion', None): 'polarizacion',
    ('voltaje ac', None): 'vac',
    ('resistencia', None): 'resistencia',
    ('ir on-off', None): 'ir_on_off',
    ('latitud', None): 'lat',
    ('longitud', None): 'lon',
    ('observaciones', None): 'observaciones',
}
_TEXTO = {'fecha', 'ref_geografica', 'observaciones'}


def _postes(wb):
    ws = wb[HOJA_PAP]
    enc = _fila_con(ws, 'abscisado', col=2)
    if enc is None:
        raise ValueError(f"'{HOJA_PAP}' sin encabezado ABSCISADO")
    mapa = _mapa(ws, enc, _COLUMNAS)
    if 'abscisa' not in mapa.values() or 'off_mv' not in mapa.values():
        raise ValueError(f"'{HOJA_PAP}': no se encontraron ABSCISADO y OFF")
    out = []
    for fila in ws.iter_rows(min_row=enc + 2, values_only=True):
        p = {}
        for i, clave in mapa.items():
            v = fila[i] if i < len(fila) else None
            if clave == 'fecha':
                p[clave] = _fecha_txt(v)
            elif clave in _TEXTO:
                p[clave] = str(v).strip() if v not in (None, '') else ''
            else:
                p[clave] = _num(v)
        if p.get('abscisa') is None:
            break          # fila numerada vacía de la plantilla, o firmas
        p['abscisa'] = int(round(p['abscisa']))
        out.append(p)
    return out


def leer_informe_pap(ruta):
    """{info, potenciales, hallazgos, fuente}."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    try:
        if HOJA_PAP not in wb.sheetnames:
            raise ValueError(f"{os.path.basename(ruta)}: no tiene la hoja "
                             f"'{HOJA_PAP}' (¿es un informe PAP?)")
        info = _info_dcvg(wb)
        info['tipo_inspeccion'] = 'PAP'
        return {'info': info, 'potenciales': _postes(wb),
                'hallazgos': _hallazgos(wb),
                'fuente': os.path.basename(ruta)}
    finally:
        wb.close()
